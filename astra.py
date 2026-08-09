import os
import logging
import time
import sqlite3
import csv

os.environ["KIVY_NO_CONSOLELOG"] = "1"
logging.getLogger("kivy").disabled = True
logging.getLogger("kivy").setLevel(logging.CRITICAL)

try:
    import openai
    HAS_OPENAI = True
except ImportError:
    openai = None
    HAS_OPENAI = False

try:
    import psutil
    HAS_PSUTIL = True
except ImportError:
    psutil = None
    HAS_PSUTIL = False

try:
    import PyPDF2
    HAS_PYPDF2 = True
except ImportError:
    PyPDF2 = None
    HAS_PYPDF2 = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    docx = None
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    openpyxl = None
    HAS_OPENPYXL = False

try:
    import speech_recognition as sr
    HAS_SR = True
except ImportError:
    sr = None
    HAS_SR = False

try:
    import pyttsx3
    HAS_PYTTSX3 = True
except ImportError:
    pyttsx3 = None
    HAS_PYTTSX3 = False

try:
    import pyautogui
    HAS_PYAUTOGUI = True
except ImportError:
    pyautogui = None
    HAS_PYAUTOGUI = False

try:
    import mss
    HAS_MSS = True
except ImportError:
    mss = None
    HAS_MSS = False

try:
    import pytesseract
    HAS_TESSERACT = True
except ImportError:
    pytesseract = None
    HAS_TESSERACT = False

try:
    import pygetwindow as gw
    HAS_PYGETWINDOW = True
except ImportError:
    gw = None
    HAS_PYGETWINDOW = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    Image = None
    HAS_PIL = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    pd = None
    HAS_PANDAS = False

try:
    import reportlab.lib.pagesizes as psizes
    from reportlab.pdfgen import canvas
    HAS_REPORTLAB = True
except ImportError:
    psizes = None
    canvas = None
    HAS_REPORTLAB = False

try:
    import matplotlib.pyplot as plt
    HAS_MPL = True
except ImportError:
    plt = None
    HAS_MPL = False

try:
    import seaborn as sns
    HAS_SNS = True
except ImportError:
    sns = None
    HAS_SNS = False

try:
    import requests
    from bs4 import BeautifulSoup
    HAS_REQUESTS = True
except ImportError:
    requests = None
    BeautifulSoup = None
    HAS_REQUESTS = False

try:
    from cryptography.fernet import Fernet
    HAS_CRYPTOGRAPHY = True
except ImportError:
    Fernet = None
    HAS_CRYPTOGRAPHY = False

try:
    import schedule
    HAS_SCHEDULE = True
except ImportError:
    schedule = None
    HAS_SCHEDULE = False

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    pdfplumber = None
    HAS_PDFPLUMBER = False

try:
    import pydub
    HAS_PYDUB = True
except ImportError:
    pydub = None
    HAS_PYDUB = False

try:
    import skimage
    HAS_SKIMAGE = True
except ImportError:
    skimage = None
    HAS_SKIMAGE = False

try:
    import imageio
    HAS_IMAGEIO = True
except ImportError:
    imageio = None
    HAS_IMAGEIO = False

try:
    import librosa
    HAS_LIBROSA = True
except ImportError:
    librosa = None
    HAS_LIBROSA = False

try:
    import sounddevice as sd
    HAS_SD = True
except ImportError:
    sd = None
    HAS_SD = False

try:
    import torch
    HAS_TORCH = True
except ImportError:
    torch = None
    HAS_TORCH = False

try:
    import tensorflow as tf
    HAS_TF = True
except ImportError:
    tf = None
    HAS_TF = False

try:
    import keras
    HAS_KERAS = True
except ImportError:
    keras = None
    HAS_KERAS = False

try:
    import httpx
    HAS_HTTPX = True
except ImportError:
    httpx = None
    HAS_HTTPX = False

try:
    import aiohttp
    HAS_AIOHTTP = True
except ImportError:
    aiohttp = None
    HAS_AIOHTTP = False

try:
    import websocket
    HAS_WEBSOCKET = True
except ImportError:
    websocket = None
    HAS_WEBSOCKET = False

try:
    import flask
    HAS_FLASK = True
except ImportError:
    flask = None
    HAS_FLASK = False

try:
    import fastapi
    HAS_FASTAPI = True
except ImportError:
    fastapi = None
    HAS_FASTAPI = False

try:
    import bcrypt
    HAS_BCRYPT = True
except ImportError:
    bcrypt = None
    HAS_BCRYPT = False

try:
    import paramiko
    HAS_PARAMIKO = True
except ImportError:
    paramiko = None
    HAS_PARAMIKO = False

try:
    import jwt
    HAS_JWT = True
except ImportError:
    jwt = None
    HAS_JWT = False

try:
    import keyring
    HAS_KEYRING = True
except ImportError:
    keyring = None
    HAS_KEYRING = False

try:
    from colorama import Fore, Style
    HAS_COLORAMA = True
except ImportError:
    class _ForeStub:
        GREEN = CYAN = YELLOW = RED = WHITE = ""
    class _StyleStub:
        RESET_ALL = ""
    Fore = _ForeStub()
    Style = _StyleStub()
    HAS_COLORAMA = False

try:
    from rich.console import Console
    _rich_console = Console()
    HAS_RICH = True
except ImportError:
    _rich_console = None
    HAS_RICH = False

try:
    from tabulate import tabulate
    HAS_TABULATE = True
except ImportError:
    tabulate = None
    HAS_TABULATE = False

try:
    import yaml
    HAS_YAML = True
except ImportError:
    yaml = None
    HAS_YAML = False

try:
    import h5py
    HAS_H5PY = True
except ImportError:
    h5py = None
    HAS_H5PY = False

try:
    import PyQt5
    HAS_PYQT5 = True
except ImportError:
    PyQt5 = None
    HAS_PYQT5 = False

try:
    import kivy
    HAS_KIVY = True
except ImportError:
    kivy = None
    HAS_KIVY = False

try:
    import dearpygui.dearpygui as dpg
    HAS_DPG = True
except ImportError:
    dpg = None
    HAS_DPG = False

try:
    import keyboard
    HAS_KEYBOARD = True
except Exception:
    keyboard = None
    HAS_KEYBOARD = False

try:
    import mouse
    HAS_MOUSE = True
except Exception:
    mouse = None
    HAS_MOUSE = False

try:
    import pytube
    HAS_PYTUBE = True
except ImportError:
    pytube = None
    HAS_PYTUBE = False

try:
    from deep_translator import GoogleTranslator
    HAS_DEEP_TRANSLATOR = True
except ImportError:
    GoogleTranslator = None
    HAS_DEEP_TRANSLATOR = False


if HAS_OPENAI:
    openai.api_key = os.getenv("OPENAI_API_KEY")


def init_db():
    conn = sqlite3.connect("memoria.db")
    c = conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS memoria (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_input TEXT,
                    ai_response TEXT,
                    timestamp TEXT
                )""")
    conn.commit()
    conn.close()


def guardar_memoria(user_input, ai_response):
    conn = sqlite3.connect("memoria.db")
    c = conn.cursor()
    c.execute("INSERT INTO memoria (user_input, ai_response, timestamp) VALUES (?, ?, ?)",
              (user_input, ai_response, time.ctime()))
    conn.commit()
    conn.close()


def cargar_memoria(limit=10):
    conn = sqlite3.connect("memoria.db")
    c = conn.cursor()
    c.execute("SELECT user_input, ai_response FROM memoria ORDER BY id DESC LIMIT ?", (limit,))
    rows = c.fetchall()
    conn.close()
    memoria = ""
    for u, a in rows[::-1]:
        memoria += f"Tú: {u}\nCopilot: {a}\n"
    return memoria


def system_status():
    if not HAS_PSUTIL:
        return "Estado del sistema no disponible."
    cpu = psutil.cpu_percent(interval=1)
    ram = psutil.virtual_memory().percent
    return f"CPU: {cpu}% | RAM: {ram}%"


def leer_pdf(ruta_pdf):
    if not HAS_PYPDF2:
        return "PyPDF2 no disponible."
    try:
        with open(ruta_pdf, "rb") as f:
            lector = PyPDF2.PdfReader(f)
            texto = ""
            for pagina in lector.pages:
                texto += pagina.extract_text() + "\n"
            return texto[:1000]
    except Exception as e:
        return f"Error al leer PDF: {e}"


def leer_word(ruta_docx):
    if not HAS_DOCX:
        return "python-docx no disponible."
    try:
        doc = docx.Document(ruta_docx)
        texto = "\n".join([p.text for p in doc.paragraphs])
        return texto[:1000]
    except Exception as e:
        return f"Error al leer Word: {e}"


def leer_excel(ruta_xlsx):
    if not HAS_OPENPYXL:
        return "openpyxl no disponible."
    try:
        wb = openpyxl.load_workbook(ruta_xlsx)
        hoja = wb.active
        texto = ""
        for fila in hoja.iter_rows(values_only=True):
            texto += " | ".join([str(c) for c in fila if c is not None]) + "\n"
        return texto[:1000]
    except Exception as e:
        return f"Error al leer Excel: {e}"


def leer_csv(ruta_csv, analizar=False):
    if not HAS_PANDAS:
        return "pandas no disponible."
    try:
        df = pd.read_csv(ruta_csv)
        if analizar:
            resumen = f"Columnas: {list(df.columns)}\nFilas totales: {len(df)}\n\nEstadísticas:\n"
            resumen += df.describe(include="all").to_string()
            return resumen[:1500]
        else:
            return df.head(10).to_string()
    except Exception as e:
        return f"Error al leer CSV: {e}"


def escribe_pdf(ruta, texto):
    if not HAS_REPORTLAB:
        return "ReportLab no disponible."
    c = canvas.Canvas(ruta, pagesize=psizes.A4)
    c.drawString(100, 750, texto)
    c.save()
    return f"PDF creado en {ruta}"


def escribe_word(ruta, texto):
    if not HAS_DOCX:
        return "python-docx no disponible."
    doc = docx.Document()
    doc.add_paragraph(texto)
    doc.save(ruta)
    return f"Word creado en {ruta}"


def escribe_excel(ruta, datos):
    if not HAS_OPENPYXL:
        return "openpyxl no disponible."
    wb = openpyxl.Workbook()
    hoja = wb.active
    for fila in datos.split("\n"):
        hoja.append(fila.split(","))
    wb.save(ruta)
    return f"Excel creado en {ruta}"


def escribe_csv(ruta, datos):
    with open(ruta, "w", newline="", encoding="utf-8") as f:
        f.write(datos)
    return f"CSV creado en {ruta}"


def grafica_csv(ruta):
    if not HAS_PANDAS or not HAS_MPL:
        return "pandas/matplotlib no disponibles."
    try:
        df = pd.read_csv(ruta)
        plt.figure(figsize=(8, 6))
        if HAS_SNS:
            sns.histplot(df[df.columns[0]], kde=True)
        else:
            plt.hist(df[df.columns[0]].dropna())
        plt.savefig("grafico.png")
        plt.close()
        return "Gráfico guardado como grafico.png"
    except Exception as e:
        return f"Error al graficar: {e}"


def extrae_web(url):
    if not HAS_REQUESTS:
        return "requests no disponible."
    try:
        r = requests.get(url, timeout=10)
        soup = BeautifulSoup(r.text, "html.parser")
        return soup.get_text()[:1000]
    except Exception as e:
        return f"Error al extraer web: {e}"


def cifra_archivo(ruta):
    if not HAS_CRYPTOGRAPHY:
        return "cryptography no disponible."
    try:
        key = Fernet.generate_key()
        fernet = Fernet(key)
        with open(ruta, "rb") as f:
            data = f.read()
        cifrado = fernet.encrypt(data)
        with open(ruta + ".cifrado", "wb") as f:
            f.write(cifrado)
        return f"Archivo cifrado en {ruta}.cifrado\nClave: {key.decode()}"
    except Exception as e:
        return f"Error al cifrar archivo: {e}"


def ask_openai(user_message):
    if not HAS_OPENAI:
        return "openai no disponible. Configure OPENAI_API_KEY e instale: pip install openai"
    try:
        memoria = cargar_memoria(limit=10)
        response = openai.ChatCompletion.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "You are Copilot integrated in a local wrapper."},
                {"role": "user", "content": f"Memoria previa:\n{memoria}\n\nNueva entrada:\n{user_message}\n\nEstado del PC: {system_status()}"}
            ]
        )
        reply = response.choices[0].message.content
        guardar_memoria(user_message, reply)
        return reply
    except Exception as e:
        return f"Error al comunicarse con OpenAI: {e}"


if __name__ == "__main__":
    init_db()
    print(Fore.GREEN + "=== ASTRA — Sistema AI Modular ===" + Style.RESET_ALL)
    while True:
        try:
            user_input = input(Fore.CYAN + "Tú: " + Style.RESET_ALL).strip()
        except (EOFError, KeyboardInterrupt):
            print("\nSaliendo...")
            break

        if user_input.lower() in ["salir", "exit", "quit"]:
            break
        elif user_input.startswith("estado pc"):
            print(system_status())
        elif user_input.startswith("extrae web "):
            print(extrae_web(user_input[11:]))
        else:
            print(ask_openai(user_input))
