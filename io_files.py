import PyPDF2
import docx
import openpyxl
import pandas as pd

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    pdfplumber = None
    HAS_PDFPLUMBER = False


def crea_py(ruta, tema):
    try:
        from ai_models import ask_openai
        codigo = ask_openai(f"Genera un script en Python sobre: {tema}")
        with open(ruta, "w", encoding="utf-8") as f:
            f.write(codigo)
        return f"Archivo Python creado en {ruta} con el tema: {tema}"
    except Exception as e:
        return f"Error creando archivo Python: {e}"


def leer_pdf(ruta_pdf, usar_plumber=False):
    try:
        texto = ""
        if usar_plumber and HAS_PDFPLUMBER:
            with pdfplumber.open(ruta_pdf) as pdf:
                for pagina in pdf.pages:
                    texto += pagina.extract_text() or ""
        else:
            with open(ruta_pdf, "rb") as f:
                lector = PyPDF2.PdfReader(f)
                for pagina in lector.pages:
                    texto += pagina.extract_text() or ""
        return texto[:70000]
    except Exception as e:
        return f"Error al leer PDF: {e}"


def leer_word(ruta_docx):
    try:
        doc = docx.Document(ruta_docx)
        texto = "\n".join([p.text for p in doc.paragraphs])
        return texto[:70000]
    except Exception as e:
        return f"Error al leer Word: {e}"


def leer_excel(ruta_xlsx):
    try:
        wb = openpyxl.load_workbook(ruta_xlsx)
        hoja = wb.active
        texto = ""
        for fila in hoja.iter_rows(values_only=True):
            texto += " | ".join([str(c) for c in fila if c is not None]) + "\n"
        return texto[:70000]
    except Exception as e:
        return f"Error al leer Excel: {e}"


def leer_csv(ruta_csv, analizar=False):
    try:
        df = pd.read_csv(ruta_csv, sep=None, engine='python', on_bad_lines='skip')
        if analizar:
            resumen = f"Columnas: {list(df.columns)}\n"
            resumen += f"Filas totales: {len(df)}\n\n"
            resumen += "Estadísticas:\n"
            resumen += df.describe(include='all').to_string()
            return resumen
        else:
            return df.head(50).to_string()
    except Exception as e:
        return f"Error al leer CSV: {e}"


def escribe_pdf(ruta, texto):
    try:
        from reportlab.pdfgen import canvas
        import reportlab.lib.pagesizes as psizes
        c = canvas.Canvas(ruta, pagesize=psizes.A4)
        c.drawString(100, 750, texto)
        c.save()
        return f"PDF creado en {ruta}"
    except ImportError:
        return "reportlab no disponible. Instale: pip install reportlab"
    except Exception as e:
        return f"Error al crear PDF: {e}"


def escribe_word(ruta, texto):
    try:
        doc = docx.Document()
        doc.add_paragraph(texto)
        doc.save(ruta)
        return f"Word creado en {ruta}"
    except Exception as e:
        return f"Error al crear Word: {e}"


def escribe_excel(ruta, datos):
    try:
        wb = openpyxl.Workbook()
        hoja = wb.active
        for fila in datos.split("\n"):
            hoja.append(fila.split(","))
        wb.save(ruta)
        return f"Excel creado en {ruta}"
    except Exception as e:
        return f"Error al crear Excel: {e}"


def escribe_csv(ruta, datos):
    try:
        with open(ruta, "w", newline="", encoding="utf-8") as f:
            f.write(datos)
        return f"CSV creado en {ruta}"
    except Exception as e:
        return f"Error al crear CSV: {e}"
